"""
Reads propellant parameters from the 'Propellant' sheet of the Excel input
file and populates a `Propellant` dataclass.
Every field has a default value. Section header rows (no numeric value in
column B) are automatically skipped.
"""
from __future__ import annotations
from dataclasses import dataclass, fields, field
from pathlib import Path
from openpyxl import load_workbook
from CoolProp.CoolProp import PropsSI

SHEET_NAME   = "Propellant "   # note: trailing space in the Excel sheet name
NAME_COLUMN  = "A"
VALUE_COLUMN = "B"
FIRST_DATA_ROW = 2


@dataclass
class Propellant:
    # PARAMÈTRES
    of_ratio:      float = 1.5
    isp:           float = 220.0
    thrust:        float = 6500.0
    total_impulse: float = 40500.0

    # THRUST CURVE MODEL
    thrust_curve_mode: str = "constant"   # "constant" ou "peaked"
    derated_thrust:    float = 0.0        # [N] palier après le pic (mode "peaked" uniquement) ; 0 = 85% du pic par défaut

    # MASS FLOW RATE
    lox_boil_off:     float = 0.001
    lox_prechill:     float = 4.362
    lox_ignition:     float = 4.362
    lox_ramp_up:      float = 2.623168871
    lox_ramp_down:    float = 2.623168871
    lox_cutoff:       float = 3.634769409
    eth_ignition:     float = 0.0
    eth_ramp_up:      float = 1.672500989
    eth_ramp_down:    float = 1.672500989
    eth_cutoff:       float = 2.29321782
    eth_frac_cooling: float = 0.075

    # MASS

    lox_temperature:  float = 90.15  # K
    lox_pressure:     float = 6e6    # Pa

    eth_temperature:  float = 293.15 # K
    eth_pressure:     float = 6e6    # Pa    
    eth_massic_frac:  float = 0.9    # massic fraction of ETH, H2O in Mischung

    # will be calculated after
    m_lox_density:    float = field(init=False, default=1154)
    m_eth_density:    float = field(init=False, default=810)

    m_eth_unused:  float = 0.405
    m_lox_unused:  float = 0.577
    m_n2:          float = 2.605

    # TIME
    hold_time:      float = 300.0
    prechill_time:  float = 0.2
    ignition_delay: float = -0.05
    ramp_up_time:   float = 0.4
    ramp_down_time: float = 0.2
    cutoff_time:    float = 0.025

    # FORCE
    f_shutdown:  float = 0.0
    f_hold_down: float = 3300.0


    def __post_init__(self):
        # LOx
        self.m_lox_density = PropsSI('D', 'T', self.lox_temperature, 'P', self.lox_pressure, 'Oxygen')

        # eth
        h2o_massic_frac = 1 - self.eth_massic_frac

        M_eth = PropsSI('molar_mass', 'Ethanol')
        M_h2o = PropsSI('molar_mass', 'Water')

        n_eth = self.eth_massic_frac / M_eth
        n_h2o = h2o_massic_frac / M_h2o

        x_eth = n_eth / (n_eth + n_h2o)
        x_h2o = 1 - x_eth

        fuel = f"Ethanol[{x_eth}]&Water[{x_h2o}]"

        self.m_eth_density = PropsSI('D', 'T', self.eth_temperature, 'P', self.eth_pressure, fuel)


def _try_to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


def read_propellant(filepath: str | Path, sheet_name: str = SHEET_NAME) -> Propellant:
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    wb = load_workbook(filepath, data_only=True, read_only=True)
    # Try exact match first, then strip-match (handles trailing spaces)
    if sheet_name not in wb.sheetnames:
        stripped = {s.strip(): s for s in wb.sheetnames}
        if sheet_name.strip() in stripped:
            sheet_name = stripped[sheet_name.strip()]
        else:
            raise ValueError(
                f"Sheet {sheet_name!r} not found in {filepath}. "
                f"Available sheets: {wb.sheetnames}"
            )

    ws = wb[sheet_name]
    name_col_idx  = ord(NAME_COLUMN)  - ord("A") + 1
    value_col_idx = ord(VALUE_COLUMN) - ord("A") + 1

    file_values: dict[str, float | str] = {}
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        name = ws.cell(row=row, column=name_col_idx).value
        if not isinstance(name, str):
            continue
        name = name.strip()
        if not name:
            continue
        raw = ws.cell(row=row, column=value_col_idx).value

        if name == "thrust_curve_mode" and isinstance(raw, str):
            file_values[name] = raw.strip().lower()
            continue

        parsed = _try_to_float(raw)
        if parsed is not None:
            file_values[name] = parsed

    valid = {f.name for f in fields(Propellant) if f.init}
    unknown = set(file_values) - valid
    if unknown:
        print(f"Warning [Propellant]: champs ignorés (inconnus) : {unknown}")

    if file_values.get("thrust_curve_mode") not in (None, "constant", "peaked"):
        print(
            f"Warning [Propellant]: thrust_curve_mode={file_values['thrust_curve_mode']!r} "
            "invalide (attendu 'constant' ou 'peaked'), retombe sur 'constant'."
        )
        file_values["thrust_curve_mode"] = "constant"

    return Propellant(**{k: v for k, v in file_values.items() if k in valid})