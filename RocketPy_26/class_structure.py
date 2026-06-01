"""
Reads structural / geometric rocket parameters from the 'Structure' sheet
of the Excel input file and populates a `Structure` dataclass.
Every field has a default value. Rows with section headers (no numeric value
in column B) are automatically skipped.
"""
from __future__ import annotations
from dataclasses import dataclass, fields
from pathlib import Path
from openpyxl import load_workbook

SHEET_NAME   = "Structure"
NAME_COLUMN  = "A"
VALUE_COLUMN = "B"
FIRST_DATA_ROW = 2


@dataclass
class Structure:
    # ROCKET
    m_additions:  float = 0.0
    r_rocket:     float = 0.1215
    drag_coeff_rocket: float = 0.38

    # BOAT TAIL
    top_r_bt:    float = 0.1215
    bottom_r_bt: float = 0.070
    l_bt:        float = 0.41755
    m_bt:        float = 1.0
    r_nozzle:    float = 0.04946

    # ENGINE BAY
    l_ebay: float = 0.756
    m_ebay: float = 26.476

    # LOX BAY
    l_lox: float = 0.388
    m_lox: float = 8.387

    # AEROCOVER
    pos_aerocover: float = 0.55
    m_aerocover:   float = 2.41
    l_aerocover:   float = 1.6558
    r_aerocover:   float = 0.03

    # PBAY 1
    l_pbay1: float = 0.9
    m_pbay1: float = 11.049

    # ETH BAY
    l_eth: float = 0.388
    m_eth: float = 8.387

    # PBAY 2
    l_pbay2: float = 0.9
    m_pbay2: float = 11.03

    # AV BAY
    l_avbay: float = 0.35
    m_avbay: float = 9.232

    # RE BAY
    l_rebay:  float = 0.715
    m_rebay:  float = 11.044
    drag_coeff_para_reefed:   float = 1.6
    drag_coeff_para_unreefed: float = 14.06
    trig_alt_para: float = 400.0

    # NOSE CONE
    l_nosecone: float = 1.003
    m_nosecone: float = 5.838

    # TANK
    r_int: float = 0.115
    h_cyl: float = 0.38510
    h_cap: float = 0.05

    # COPV
    r_int_copv: float = 0.0885
    h_copv:     float = 0.570

    # FINS
    n_fins:       float = 4.0
    env_fins:     float = 0.3
    rc_fins:      float = 0.671
    tc_fins:      float = 0.161
    pos_fins:     float = 1.1315
    cangle_fins:  float = 0.0
    sweep_l_fins: float = 0.539

    # RAIL BUTTON
    pos_sup_rb: float = 0.46280
    pos_inf_rb: float = 0.46279


def _try_to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


def read_structure(filepath: str | Path, sheet_name: str = SHEET_NAME) -> Structure:
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    wb = load_workbook(filepath, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet {sheet_name!r} not found in {filepath}. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    name_col_idx  = ord(NAME_COLUMN)  - ord("A") + 1
    value_col_idx = ord(VALUE_COLUMN) - ord("A") + 1

    file_values: dict[str, float] = {}
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        name = ws.cell(row=row, column=name_col_idx).value
        if not isinstance(name, str):
            continue
        name = name.strip()
        if not name:
            continue
        parsed = _try_to_float(ws.cell(row=row, column=value_col_idx).value)
        if parsed is not None:
            file_values[name] = parsed

    valid = {f.name for f in fields(Structure)}
    unknown = set(file_values) - valid
    if unknown:
        print(f"Warning [Structure]: champs ignorés (inconnus) : {unknown}")

    return Structure(**{k: v for k, v in file_values.items() if k in valid})
