from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

TARGET_SHEET = "Simulation"   
NAME_COLUMN = "A"             
VALUE_COLUMN = "B"            
START_ROW = 2                 

OUTPUT_STRUCTURE = [
    ("PERFORMANCES MOTEUR", "TITRE"),
    ("Max thrust [N]", "max_thrust"),
    ("Total impulse [Ns]", "total_impulse"),
    
    ("CINÉMATIQUE & TRAJECTOIRE", "TITRE"),
    ("Apogee AGL [m]", "apogee_agl"),
    ("Apogee ASL [m]", "apogee_asl"),
    ("Max velocity [m/s]", "max_speed"),
    ("Max acceleration [m/s^2]", "max_acceleration"),
    ("Max Mach [-]", "max_mach"),
    ("Time to apogee [s]", "time_to_apogee"),
    ("Total flight time [s]", "total_flight_time"),
    
    ("SORTIE DE RAMPE", "TITRE"),
    ("Rail exit velocity [m/s]", "rail_exit_vel"),
    ("Static margin rail exit [c]", "margin_rail"),
    
    ("STABILITÉ", "TITRE"),
    ("Static margin liftoff [c]", "margin_liftoff"),
]

EXTRACTORS = {
    "max_thrust":        lambda fl: float(fl.rocket.motor.max_thrust if callable(fl.rocket.motor.max_thrust) else fl.rocket.motor.max_thrust),
    "total_impulse":     lambda fl: float(fl.rocket.motor.total_impulse if callable(fl.rocket.motor.total_impulse) else fl.rocket.motor.total_impulse),
    "apogee_agl":        lambda fl: float(fl.altitude(fl.apogee_time)),
    "apogee_asl":        lambda fl: float(fl.apogee),
    "max_speed":         lambda fl: float(fl.max_speed),
    "max_acceleration":  lambda fl: float(fl.max_acceleration),
    "max_mach":          lambda fl: float(fl.max_mach_number),
    "time_to_apogee":    lambda fl: float(fl.apogee_time),
    "total_flight_time": lambda fl: float(fl.t_final),
    "rail_exit_vel":     lambda fl: float(fl.out_of_rail_velocity),
    "margin_liftoff":    lambda fl: float(fl.static_margin(0)),
    "margin_rail":       lambda fl: float(fl.static_margin(fl.out_of_rail_time)),
}

def write_flight_outputs_vertical(
    flight: Any,
    version: str,
    directory: str | Path = ".",
    dry_mass: float | None = None,
    wet_mass: float | None = None,
    nominal_total_impulse: float | None = None,
) -> dict[str, Any]:
    """Cherche l'Excel dans le dossier indiqué et écrit les résultats
    verticalement. dry_mass/wet_mass sont ajoutés en tête (section MASSE)
    s'ils sont fournis. nominal_total_impulse, si fourni, remplace la
    valeur 'Total impulse' calculée par RocketPy (qui exclut la phase
    hold-down) par l'impulsion nominale demandée (qui l'inclut)."""
    filepath = Path(directory) / f"{version}.xlsx"
    
    if not filepath.exists():
        raise FileNotFoundError(f"Le fichier Excel '{filepath.name}' est introuvable au même endroit.")

    wb = load_workbook(filepath, data_only=False)
    if TARGET_SHEET not in wb.sheetnames:
        raise ValueError(f"L'onglet '{TARGET_SHEET}' n'existe pas.")
        
    ws = wb[TARGET_SHEET]
    
    # Nettoyage
    for r in range(START_ROW, START_ROW + 40):
        for merged_range in list(ws.merged_cells.ranges):
            if r in range(merged_range.min_row, merged_range.max_row + 1):
                ws.unmerge_cells(str(merged_range))
        ws[f"{NAME_COLUMN}{r}"] = None
        ws[f"{VALUE_COLUMN}{r}"] = None
        ws[f"{NAME_COLUMN}{r}"].fill = PatternFill(fill_type=None)
        ws[f"{NAME_COLUMN}{r}"].font = Font(name="Arial", size=10, bold=False)
    
    # Styles
    titre_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    titre_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    data_font_bold = Font(name="Arial", size=10, bold=False)
    
    current_row = START_ROW
    written = {}

    if dry_mass is not None or wet_mass is not None:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        cell = ws[f"{NAME_COLUMN}{current_row}"]
        cell.value = "MASSE"
        cell.font = titre_font
        cell.fill = titre_fill
        current_row += 1
        if dry_mass is not None:
            ws[f"{NAME_COLUMN}{current_row}"] = "Dry mass [kg]"
            ws[f"{NAME_COLUMN}{current_row}"].font = data_font_bold
            ws[f"{VALUE_COLUMN}{current_row}"] = float(dry_mass)
            ws[f"{VALUE_COLUMN}{current_row}"].number_format = "0.00"
            written["Dry mass [kg]"] = float(dry_mass)
            current_row += 1
        if wet_mass is not None:
            ws[f"{NAME_COLUMN}{current_row}"] = "Wet mass (initial) [kg]"
            ws[f"{NAME_COLUMN}{current_row}"].font = data_font_bold
            ws[f"{VALUE_COLUMN}{current_row}"] = float(wet_mass)
            ws[f"{VALUE_COLUMN}{current_row}"].number_format = "0.00"
            written["Wet mass (initial) [kg]"] = float(wet_mass)
            current_row += 1

    for label, key in OUTPUT_STRUCTURE:
        if key == "TITRE":
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            cell = ws[f"{NAME_COLUMN}{current_row}"]
            cell.value = label
            cell.font = titre_font
            cell.fill = titre_fill
            current_row += 1
            continue

        if key == "total_impulse" and nominal_total_impulse is not None:
            value = float(nominal_total_impulse)
        else:
            extractor = EXTRACTORS.get(key)
            try:
                value = extractor(flight)
            except Exception:
                continue

        if value is None:
            continue

        ws[f"{NAME_COLUMN}{current_row}"] = label
        ws[f"{NAME_COLUMN}{current_row}"].font = data_font_bold
        ws[f"{VALUE_COLUMN}{current_row}"] = value
        ws[f"{VALUE_COLUMN}{current_row}"].number_format = "0.0000" if abs(value) < 10 else "0.00"
        written[label] = value
        current_row += 1  

    wb.save(filepath)
    print(f"✅ Succès ! Fichier '{filepath.name}' mis à jour au même endroit.")
    return written